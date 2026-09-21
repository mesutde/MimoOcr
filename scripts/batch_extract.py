#!/usr/bin/env python3
"""Extract plain text from non-image files for Mimo OCR batch mode.

Usage: batch_extract.py <file>
Prints UTF-8 text to stdout. Exit 0 on success.
Supported: pdf, docx, xlsx, pptx, udf (UYAP), txt, md, csv, json, rtf(simple)
"""

from __future__ import annotations

import html
import sys
import zipfile
import re
from pathlib import Path
from xml.etree import ElementTree as ET


def read_plain(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "cp1254", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError as e:
            raise RuntimeError(f"pypdf not available: {e}") from e
    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages):
        try:
            t = page.extract_text() or ""
        except Exception as e:
            t = f"[page {i + 1} error: {e}]"
        if t.strip():
            parts.append(t)
    return "\n\n".join(parts)


def extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs if p.text and p.text.strip()]
        # tables
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    paras.append(" | ".join(cells))
        return "\n".join(paras)
    except Exception:
        # Fallback: raw XML text from word/document.xml
        with zipfile.ZipFile(path, "r") as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="replace")
        texts = re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)
        return " ".join(texts)


def extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(c.strip() for c in cells):
                    lines.append("\t".join(cells))
        wb.close()
        return "\n".join(lines)
    except Exception as e:
        raise RuntimeError(f"xlsx: {e}") from e


def extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"## Slide {i}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text and shape.text.strip():
                    parts.append(shape.text)
        return "\n".join(parts)
    except Exception:
        # raw XML fallback
        with zipfile.ZipFile(path, "r") as z:
            slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
            parts = []
            for name in sorted(slides):
                xml = z.read(name).decode("utf-8", errors="replace")
                texts = re.findall(r"<a:t>([^<]*)</a:t>", xml)
                if texts:
                    parts.append(" ".join(texts))
        return "\n".join(parts)


def _strip_rtf(rtf: str) -> str:
    rtf = re.sub(r"\\'[0-9a-fA-F]{2}", lambda m: chr(int(m.group(0)[2:], 16)), rtf)
    rtf = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", rtf)
    rtf = re.sub(r"[{}]", "", rtf)
    rtf = re.sub(r"\s+", " ", rtf)
    return rtf.strip()


def extract_rtf(path: Path) -> str:
    return _strip_rtf(read_plain(path))


def _xml_local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _collect_text_xml(xml_bytes: bytes) -> str:
    parts: list[str] = []
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError:
        text = xml_bytes.decode("utf-8", errors="replace")
        chunks = re.findall(r">([^<>]+)<", text)
        return "\n".join(html.unescape(c).strip() for c in chunks if c.strip())

    para_tags = {"paragraph", "p", "textparagraph", "par", "block", "content"}
    line_tags = {"line", "br", "tab"}

    def walk(node, buf: list[str]) -> None:
        name = _xml_local(node.tag).lower()
        if node.text:
            t = node.text.strip()
            if t:
                buf.append(t)
        for child in list(node):
            cname = _xml_local(child.tag).lower()
            if cname in {"signature", "binary", "imagedata", "image"}:
                continue
            if name in para_tags:
                walk(child, buf)
                buf.append("\n")
            elif cname in line_tags:
                walk(child, buf)
                buf.append("\n")
            else:
                walk(child, buf)
            if child.tail:
                tt = child.tail.strip()
                if tt:
                    buf.append(tt)

    buf: list[str] = []
    walk(root, buf)
    text = "".join(buf)
    raw = xml_bytes.decode("utf-8", errors="replace")
    for frag in re.findall(r"\{\\rtf1(?:[^{}]|\{[^{}]*\})*\}", raw):
        st = _strip_rtf(frag)
        if st:
            text += "\n" + st
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_udf(path: Path) -> str:
    """UYAP Doküman Formatı (.udf): ZIP + content.xml (+ signature.p7s, binary/).

    Justice Ministry UYAP container (not optical-disc UDF).
    """
    if not zipfile.is_zipfile(path):
        raise RuntimeError("not a UYAP UDF zip container (or file is corrupt)")
    with zipfile.ZipFile(path, "r") as z:
        names = z.namelist()
        candidates = ["content.xml", "Content.xml", "CONTENT.XML", "content/Content.xml"]
        xml_name = next((c for c in candidates if c in names), None)
        if xml_name is None:
            xml_name = next((n for n in names if n.lower().endswith("content.xml")), None)
        if xml_name is None:
            xml_name = next(
                (n for n in names if n.lower().endswith(".xml") and "sign" not in n.lower()),
                None,
            )
        if xml_name is None:
            raise RuntimeError(f"UDF content.xml not found; entries={names[:20]}")

        text = _collect_text_xml(z.read(xml_name))
        has_sig = any(n.lower().endswith(".p7s") or "signature" in n.lower() for n in names)
        binaries = [n for n in names if n.lower().startswith("binary/")]
        meta = [f"[UYAP UDF] content={xml_name}"]
        if has_sig:
            meta.append("[e-signature present — not exported as text]")
        if binaries:
            meta.append(f"[embedded objects: {len(binaries)}]")
        body = text if text.strip() else "[content.xml has no extractable text — image-only?]"
        return "\n".join(meta) + "\n\n" + body


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: batch_extract.py <file>", file=sys.stderr)
        return 2
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"file not found: {path}", file=sys.stderr)
        return 1
    ext = path.suffix.lower().lstrip(".")
    try:
        if ext == "pdf":
            text = extract_pdf(path)
        elif ext == "docx":
            text = extract_docx(path)
        elif ext == "xlsx":
            text = extract_xlsx(path)
        elif ext == "pptx":
            text = extract_pptx(path)
        elif ext == "udf":
            text = extract_udf(path)
        elif ext == "rtf":
            text = extract_rtf(path)
        elif ext in {"txt", "md", "csv", "json", "log", "xml", "html", "htm"}:
            text = read_plain(path)
        else:
            print(f"unsupported type: .{ext}", file=sys.stderr)
            return 3
        sys.stdout.write(text)
        return 0
    except Exception as e:
        print(f"extract error: {e}", file=sys.stderr)
        return 4


if __name__ == "__main__":
    raise SystemExit(main())
